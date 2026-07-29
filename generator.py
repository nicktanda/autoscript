#!/usr/bin/env python3
"""TKE SWMS Generator v2.0 — generates site-specific SWMS packages from .docx templates."""

import customtkinter as ctk
import tkinter as tk
from tkinter import filedialog, messagebox
from datetime import datetime
from docx import Document
from docx.oxml.ns import qn
from PIL import Image, ImageDraw, ImageFont
import webbrowser
import zipfile
import subprocess
import platform
import os
import json

SCRIPT_DIR = os.path.dirname(os.path.abspath(__file__))
GRADIENT_IMG = os.path.join(SCRIPT_DIR, "tke_logo_gradient_reveal.png")

CONFIG_FILE = os.path.join(os.path.dirname(os.path.abspath(__file__)), "config.json")

SHAREPOINT_URLS = {
    "NI": "https://tke.sharepoint.com/:f:/s/tkeauteamsites/qld/IgA5kddyxW9wRbINLsNj43JhAdiHowU2WXFqFBJrXadWvtY",
    "MOD": "https://tke.sharepoint.com/:f:/s/tkeauteamsites/qld/IgBB6FCXD-6KTL8P0jvc19hUAdRWxr96_Y4uVUrYISxKyy8",
    "RENEW": "https://tke.sharepoint.com/:f:/s/tkeauteamsites/qld/IgAqXfBHcI1qQKS0GMMO2ZEJAR8lO9Q5L3rBqu2zhyyuLC0",
}

JOB_FILTERS = {
    "NI": lambda f: f.startswith("NI"),
    "MOD": lambda f: f.startswith("M") and not f.startswith("NI"),
    "RENEW": lambda f: f.startswith("R"),
}


# ── Config ──────────────────────────────────────────────────────────────


def load_config():
    if os.path.exists(CONFIG_FILE):
        try:
            with open(CONFIG_FILE, "r") as f:
                return json.load(f)
        except Exception:
            pass
    return {}


def save_config(data):
    with open(CONFIG_FILE, "w") as f:
        json.dump(data, f, indent=4)


# ── Document processing (image-safe) ───────────────────────────────────


def _replace_in_paragraph_element(para_elem, old_text, new_text):
    """Replace old_text→new_text across w:t elements in a single <w:p>, handling
    the case where Word's proofing engine splits the placeholder across multiple
    XML runs (e.g. '<job name>' → ['<','j','ob name>'])."""

    t_elems = list(para_elem.iter(qn("w:t")))
    if not t_elems:
        return

    texts = [t.text or "" for t in t_elems]

    # ── Pass 1: simple single-element replacement ──
    for i, txt in enumerate(texts):
        if old_text in txt:
            texts[i] = txt.replace(old_text, new_text)
            t_elems[i].text = texts[i]
            # Preserve whitespace
            t_elems[i].set(qn("xml:space"), "preserve")

    # ── Pass 2: cross-element replacement ──
    full = "".join(texts)
    while old_text in full:
        idx = full.find(old_text)
        end = idx + len(old_text)

        # Map character offsets → element indices
        cum = 0
        ranges = []
        for t in texts:
            ranges.append((cum, cum + len(t)))
            cum += len(t)

        first = last = None
        for i, (s, e) in enumerate(ranges):
            if first is None and e > idx:
                first = i
            if e >= end:
                last = i
                break

        if first is None or last is None:
            break

        prefix = texts[first][: idx - ranges[first][0]]
        suffix = texts[last][end - ranges[last][0] :]

        # Stitch: first element gets prefix+replacement, last gets suffix, middle cleared
        if first == last:
            texts[first] = prefix + new_text + suffix
        else:
            texts[first] = prefix + new_text
            for m in range(first + 1, last):
                texts[m] = ""
            texts[last] = suffix

        # Write back
        for i in range(first, last + 1):
            t_elems[i].text = texts[i]
            t_elems[i].set(qn("xml:space"), "preserve")

        full = "".join(texts)


def _replace_in_tree(root_element, old_text, new_text):
    """Walk every <w:p> inside root_element and replace placeholders."""
    for para in root_element.iter(qn("w:p")):
        _replace_in_paragraph_element(para, old_text, new_text)


def process_document(input_path, output_path, site, date_text):
    """Open a .docx, replace placeholders while preserving images & formatting, save."""
    doc = Document(input_path)

    replacements = [("<job name>", site), ("<0/0/0>", date_text)]

    for old, new in replacements:
        # Main document body
        _replace_in_tree(doc.element, old, new)

        # Headers & footers (separate XML parts)
        for rel in doc.part.rels.values():
            if "header" in rel.reltype or "footer" in rel.reltype:
                _replace_in_tree(rel.target_part.element, old, new)

    doc.save(output_path)


def open_folder(path):
    system = platform.system()
    if system == "Darwin":
        subprocess.Popen(["open", path])
    elif system == "Windows":
        os.startfile(path)
    else:
        subprocess.Popen(["xdg-open", path])


# ── UI ──────────────────────────────────────────────────────────────────


def _make_gradient_title(text, font_size=40, bg_hex="#242424"):
    """Render *text* with the TKE gradient image clipped to the letter shapes."""
    gradient = Image.open(GRADIENT_IMG).convert("RGBA")

    # Try to find a bold system font; fall back to default
    try:
        font = ImageFont.truetype("Arial Bold.ttf", font_size)
    except OSError:
        for name in ("ArialB.ttf", "Helvetica-Bold.ttf", "DejaVuSans-Bold.ttf",
                      "/System/Library/Fonts/Helvetica.ttc"):
            try:
                font = ImageFont.truetype(name, font_size)
                break
            except OSError:
                continue
        else:
            font = ImageFont.load_default()

    # Measure text
    tmp = ImageDraw.Draw(Image.new("RGBA", (1, 1)))
    bbox = tmp.textbbox((0, 0), text, font=font)
    tw, th = bbox[2] - bbox[0], bbox[3] - bbox[1]
    pad_x, pad_y = 20, 12
    w, h = tw + pad_x * 2, th + pad_y * 2

    # Scale gradient to cover the text area
    grad = gradient.resize((w, h), Image.LANCZOS)

    # Draw text mask (white text on black)
    mask_img = Image.new("L", (w, h), 0)
    draw = ImageDraw.Draw(mask_img)
    draw.text((pad_x - bbox[0], pad_y - bbox[1]), text, fill=255, font=font)

    # Composite: gradient where mask is white, bg_hex elsewhere
    bg_r, bg_g, bg_b = int(bg_hex[1:3], 16), int(bg_hex[3:5], 16), int(bg_hex[5:7], 16)
    result = Image.new("RGBA", (w, h), (bg_r, bg_g, bg_b, 255))
    result.paste(grad, mask=mask_img)

    return result


class App(ctk.CTk):
    def __init__(self):
        super().__init__()

        self.title("TKE SWMS Generator v2.0")
        self.geometry("920x800")
        self.minsize(720, 620)

        ctk.set_appearance_mode("dark")
        ctk.set_default_color_theme("blue")

        cfg = load_config()
        downloads = os.path.join(os.path.expanduser("~"), "Downloads")

        self._working = ctk.StringVar(value=cfg.get("working_folder", downloads))
        self._output = ctk.StringVar(value=cfg.get("output_folder", downloads))
        self._job = ctk.StringVar(value="NI")
        self._site = ctk.StringVar()
        self._date = ctk.StringVar(value=datetime.today().strftime("%d/%m/%Y"))

        self._build()
        self._scan()

    # ── Layout ─────────────────────────────────────────────────────

    def _build(self):
        root = ctk.CTkFrame(self, fg_color="transparent")
        root.pack(fill="both", expand=True, padx=24, pady=20)

        # Title — gradient-clipped text
        title_pil = _make_gradient_title("TKE SWMS Generator", font_size=40)
        self._title_img = ctk.CTkImage(
            light_image=title_pil, dark_image=title_pil,
            size=(title_pil.width, title_pil.height),
        )
        ctk.CTkLabel(root, image=self._title_img, text="").pack(pady=(0, 16))

        # ── Job type ─────────────────────────────────────────────
        card = ctk.CTkFrame(root)
        card.pack(fill="x", pady=(0, 8))

        ctk.CTkLabel(
            card, text="Job Type",
            font=ctk.CTkFont(size=13, weight="bold"),
        ).pack(anchor="w", padx=16, pady=(12, 4))

        row = ctk.CTkFrame(card, fg_color="transparent")
        row.pack(fill="x", padx=16, pady=(0, 10))

        for v in ("NI", "MOD", "RENEW"):
            ctk.CTkRadioButton(
                row, text=v, variable=self._job, value=v,
                command=self._scan,
            ).pack(side="left", padx=(0, 24))

        ctk.CTkButton(
            row, text="Open SharePoint", width=160,
            command=self._open_sp,
        ).pack(side="right")

        # ── Source folder ────────────────────────────────────────
        card = ctk.CTkFrame(root)
        card.pack(fill="x", pady=(0, 8))

        ctk.CTkLabel(
            card, text="Source Folder  (downloaded SWMS templates)",
            font=ctk.CTkFont(size=13, weight="bold"),
        ).pack(anchor="w", padx=16, pady=(12, 4))

        row = ctk.CTkFrame(card, fg_color="transparent")
        row.pack(fill="x", padx=16, pady=(0, 10))

        ctk.CTkEntry(row, textvariable=self._working).pack(
            side="left", fill="x", expand=True, padx=(0, 8),
        )
        ctk.CTkButton(row, text="Browse", width=90, command=self._browse_src).pack(side="right")

        # ── Document list ────────────────────────────────────────
        card = ctk.CTkFrame(root)
        card.pack(fill="both", expand=True, pady=(0, 8))

        hdr = ctk.CTkFrame(card, fg_color="transparent")
        hdr.pack(fill="x", padx=16, pady=(12, 4))

        ctk.CTkLabel(
            hdr, text="Documents",
            font=ctk.CTkFont(size=13, weight="bold"),
        ).pack(side="left")

        self._status = ctk.CTkLabel(hdr, text="", text_color="gray")
        self._status.pack(side="right")

        lf = ctk.CTkFrame(card, fg_color="transparent")
        lf.pack(fill="both", expand=True, padx=16, pady=(0, 6))

        self._lb = tk.Listbox(
            lf, selectmode=tk.MULTIPLE,
            bg="#2b2b2b", fg="#dcdcdc",
            selectbackground="#1f6aa5", selectforeground="#ffffff",
            highlightthickness=0, bd=0, relief="flat",
            font=("Segoe UI", 11), activestyle="none",
        )
        sb = ctk.CTkScrollbar(lf, command=self._lb.yview)
        self._lb.configure(yscrollcommand=sb.set)
        self._lb.pack(side="left", fill="both", expand=True)
        sb.pack(side="right", fill="y")

        brow = ctk.CTkFrame(card, fg_color="transparent")
        brow.pack(fill="x", padx=16, pady=(0, 10))

        ctk.CTkButton(brow, text="Select All", width=100, command=self._sel_all).pack(side="left", padx=(0, 8))
        ctk.CTkButton(brow, text="Select None", width=100, command=self._sel_none).pack(side="left")

        # ── Site / Date / Output ─────────────────────────────────
        card = ctk.CTkFrame(root)
        card.pack(fill="x", pady=(0, 8))

        inner = ctk.CTkFrame(card, fg_color="transparent")
        inner.pack(fill="x", padx=16, pady=12)
        inner.columnconfigure(0, weight=1)

        ctk.CTkLabel(inner, text="Site Name", font=ctk.CTkFont(weight="bold")).grid(row=0, column=0, sticky="w")
        ctk.CTkLabel(inner, text="Date", font=ctk.CTkFont(weight="bold")).grid(row=0, column=1, sticky="w", padx=(16, 0))

        ctk.CTkEntry(inner, textvariable=self._site, width=400).grid(row=1, column=0, sticky="w", pady=(2, 10))
        ctk.CTkEntry(inner, textvariable=self._date, width=140).grid(row=1, column=1, sticky="w", padx=(16, 0), pady=(2, 10))

        ctk.CTkLabel(inner, text="Output Folder", font=ctk.CTkFont(weight="bold")).grid(row=2, column=0, columnspan=2, sticky="w")

        orow = ctk.CTkFrame(inner, fg_color="transparent")
        orow.grid(row=3, column=0, columnspan=2, sticky="ew", pady=(2, 0))

        ctk.CTkEntry(orow, textvariable=self._output).pack(side="left", fill="x", expand=True, padx=(0, 8))
        ctk.CTkButton(orow, text="Browse", width=90, command=self._browse_out).pack(side="right")

        # ── Generate ─────────────────────────────────────────────
        ctk.CTkButton(
            root, text="Generate SWMS Package", command=self._generate,
            height=48, font=ctk.CTkFont(size=15, weight="bold"),
            fg_color="#2d8f2d", hover_color="#237a23",
        ).pack(fill="x", pady=(4, 0))

    # ── Actions ────────────────────────────────────────────────────

    def _persist(self):
        save_config({
            "working_folder": self._working.get(),
            "output_folder": self._output.get(),
        })

    def _open_sp(self):
        webbrowser.open(SHAREPOINT_URLS[self._job.get()])

    def _browse_src(self):
        d = filedialog.askdirectory(title="Select Source Folder")
        if d:
            self._working.set(d)
            self._persist()
            self._scan()

    def _browse_out(self):
        d = filedialog.askdirectory(title="Select Output Folder")
        if d:
            self._output.set(d)
            self._persist()

    def _scan(self):
        self._lb.delete(0, tk.END)
        folder = self._working.get()
        if not os.path.isdir(folder):
            self._status.configure(text="Folder not found")
            return

        job = self._job.get()
        match = JOB_FILTERS[job]

        for name in sorted(os.listdir(folder)):
            if not name.lower().endswith(".docx") or name.startswith("~$"):
                continue
            if match(name.upper()):
                self._lb.insert(tk.END, name)

        self._lb.select_set(0, tk.END)
        self._status.configure(text=f"{self._lb.size()} {job} doc(s)")

    def _sel_all(self):
        self._lb.select_set(0, tk.END)

    def _sel_none(self):
        self._lb.selection_clear(0, tk.END)

    def _generate(self):
        indices = self._lb.curselection()
        if not indices:
            messagebox.showerror("Error", "Select at least one document.")
            return

        site = self._site.get().strip()
        if not site:
            messagebox.showerror("Error", "Enter a site name.")
            return

        date_text = self._date.get()
        src_folder = self._working.get()
        out_dir = os.path.join(self._output.get(), f"{site} SWMS")
        os.makedirs(out_dir, exist_ok=True)

        ok, errors = 0, []

        for i in indices:
            fname = self._lb.get(i)
            src = os.path.join(src_folder, fname)
            base, ext = os.path.splitext(fname)
            dst = os.path.join(out_dir, f"{base} - {site}{ext}")
            try:
                process_document(src, dst, site, date_text)
                ok += 1
            except Exception as exc:
                errors.append(f"{fname}: {exc}")

        # Zip all generated docs
        zip_path = os.path.join(out_dir, f"{site} SWMS.zip")
        with zipfile.ZipFile(zip_path, "w", zipfile.ZIP_DEFLATED) as zf:
            for f in os.listdir(out_dir):
                fp = os.path.join(out_dir, f)
                if fp != zip_path:
                    zf.write(fp, arcname=f)

        msg = f"Processed {ok} file(s)\nZip saved to:\n{zip_path}"
        if errors:
            msg += "\n\nErrors:\n" + "\n".join(errors)

        messagebox.showinfo("Complete", msg)

        try:
            open_folder(out_dir)
        except Exception:
            pass


if __name__ == "__main__":
    App().mainloop()
